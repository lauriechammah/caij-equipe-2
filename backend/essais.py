"""Essais hors ligne : aucun appel à Augure ni au Compagnon, aucune clé requise.

    python -m backend.essais

Vérifie la relecture JSON, la source du document, la manipulation Word et
l'assemblage des trois méthodes. Utile avant une démonstration.
"""

import json
import sys
import tempfile
import zipfile
from pathlib import Path

from docx import Document

from backend import db
from backend.acte import appliquer_plan, create_acte_modifie, texte_document
from backend.analyse import _normaliser_contradictions, get_contradictions, get_type_consequences
from backend.llm import extraire_json

_reussites, _echecs = 0, 0


def verifier(condition, libelle):
    global _reussites, _echecs
    if condition:
        _reussites += 1
        print(f"  ok   {libelle}")
    else:
        _echecs += 1
        print(f"  ÉCHEC {libelle}")


def essai_json():
    print("\nextraire_json")
    verifier(extraire_json('{"a": 1}') == {"a": 1}, "objet simple")
    verifier(extraire_json('Voici la réponse :\n```json\n{"a": [1, 2]}\n```\n') == {"a": [1, 2]},
             "réponse entourée d'un bloc de code")
    verifier(extraire_json('{"a": 1, "b": 2,}') == {"a": 1, "b": 2}, "virgule finale")
    verifier(extraire_json('pas du json') is None, "réponse sans JSON")


def essai_document():
    print("\nsource de l'acte (db)")
    dossiers = db.lister_dossiers()
    verifier(bool(dossiers), f"modèles disponibles : {', '.join(dossiers) or 'aucun'}")
    for dossier in dossiers:
        chemin = db.get_acte_source(dossier)
        verifier(chemin.exists() and chemin.suffix == ".docx", f"{dossier} : {chemin.name}")
    with tempfile.TemporaryDirectory() as dossier_temporaire:
        vide = Path(dossier_temporaire) / "vide"
        vide.mkdir()
        (vide / "x.docx").write_bytes(b"")
        modeles, base = db.MODELES, db.BASE
        db.MODELES, db.BASE = Path(dossier_temporaire), Path(dossier_temporaire) / "vide.db"
        try:
            verifier(db.lister_dossiers() == ["vide"], "un seul modèle : le dossier est déduit")
            try:
                db.get_acte_source("introuvable")
                verifier(False, "dossier inconnu : doit lever une erreur")
            except db.ErreurDocument:
                verifier(True, "dossier inconnu : erreur explicite")
        finally:
            db.MODELES, db.BASE = modeles, base


def essai_word():
    print("\nmanipulation Word")
    source = Path(tempfile.mkdtemp()) / "modele.docx"
    document = Document()
    document.add_paragraph("MANDAT DE PROTECTION")
    document.add_paragraph("La mandante désigne le mandataire pour exercer ses droits en son nom. "
                           "Le mandataire agit dans l'intérêt de la mandante.")
    table = document.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "Adresse du mandataire"
    document.add_paragraph("Fait à Québec, ce jour.")
    document.save(str(source))

    plan = {
        "surlignages": [
            {"ancre": "La mandante désigne le mandataire pour exercer ses droits en son nom.",
             "note": "Deux versions de l'identité dans le dossier.", "origine": "contradiction",
             "couleur": "jaune"},
            {"ancre": "phrase absente", "note": "Rien à faire."},
        ],
        "ajouts": [
            {"apres": "Le mandataire agit dans l'intérêt de la mandante.",
             "texte": "Toute décision doit être approuvée par le tribunal.",
             "note": "Ajout imposé par l'état de santé décrit au dossier."},
        ],
        "signalees": [
            {"clause": "Adresse du mandataire", "raison": "Le domicile n'est pas établi."},
        ],
        "questions_au_client": ["Quelle est la date de la signature ?"],
        "consequences_a_rappeler": [{"type": "nullité", "consequence": "Mandat nul si inaptitude.",
                                     "source": {"titre": "CCQ-3115", "url": "https://canlii.example/3115"}}],
    }
    modifie = Document(str(source))
    rapport = appliquer_plan(modifie, plan)
    sortie = source.with_name("acte.docx")
    modifie.save(str(sortie))

    verifier(rapport["compteurs"] == {"appliquees": 3, "introuvables": 1, "questions": 1},
             f"compteurs {rapport['compteurs']}")
    verifier(rapport["introuvables"][0]["type"] == "annotation", "l'ancre absente est signalée")
    verifier(any(o["type"] == "signalement" for o in rapport["appliquees"]),
             "la clause non applicable est signalée, pas supprimée")

    relu = Document(str(sortie))
    verifier(len(relu.comments) == 3, f"commentaires Word : {len(relu.comments)}")
    texte = texte_document(sortie)
    verifier("Toute décision doit être approuvée par le tribunal." in texte,
             "la clause ajoutée est dans le texte")
    verifier("[Note : " in texte, "la note est posée à côté du passage")
    verifier("Fait à Québec, ce jour." in texte, "le texte du modèle n'a pas été modifié")
    verifier("Adresse du mandataire" in texte, "les cellules de tableau sont lues")
    conservees = ["MANDAT DE PROTECTION", "La mandante désigne le mandataire",
                  "Le mandataire agit dans l'intérêt de la mandante.", "Fait à Québec, ce jour."]
    verifier(all(c in texte for c in conservees), "aucune clause du modèle n'a disparu")
    xml = zipfile.ZipFile(str(sortie)).read("word/document.xml").decode("utf-8")
    verifier('w:color w:val="C00000"' in xml, "les ajouts sont en rouge")
    verifier("w:highlight" in xml, "les passages sont surlignés")
    verifier(texte_document(sortie).startswith("MANDAT DE PROTECTION"), "relecture du texte")
    verifier(texte_document(sortie).index("Adresse du mandataire")
             < texte_document(sortie).index("Fait à Québec"),
             "l'ordre du document est respecté, tableaux compris")
    verifier("Annexe" in texte and "Quelle est la date de la signature ?" in texte,
             "l'annexe reprend les questions")
    verifier("https://canlii.example/3115" in texte, "l'annexe cite l'url de la source")


def essai_analyse():
    print("\nassemblage des méthodes")
    rapport = _normaliser_contradictions({"contradictions": [{"demandes_concernees": ["deux prénoms"],
                                                               "probleme": "deux versions"}],
                                          "questions_au_notaire": ["Quel prénom ?"]})
    verifier(rapport["contradictions"][0]["id"] == "C1", "les contradictions sont numérotées")
    verifier(_normaliser_contradictions([])["aucune_contradiction"] is True,
             "un rapport vide est traité")

    import backend.analyse as analyse
    demandes = []

    def faux_augure(prompt, **options):
        demandes.append(prompt)
        if "consequences" in options.get("consignes", ""):
            return {"consequences": [{"type": "nullité",
                                      "consequence": "Mandat nul si inaptitude."}]}
        return {"contradictions": [{"demandes_concernees": ["deux versions du nom"],
                                     "probleme": "deux versions", "articles": ["CCQ-3115"],
                                     "certitude": "certaine"}],
                "questions_au_notaire": ["Quel est le bon nom ?"]}

    def faux_compagnon(prompt, **options):
        return {"nb_sources": 1, "resume": "", "sources": [
            {"type": "legislation", "titre": "CCQ-3115", "url": "https://canlii.example/3115",
             "extrait": "Un mandat de protection est valide…", "score": 0.9, "rang": 1,
             "metadonnees": {}}]}

    analyse.ask_augure, analyse.ask_compagnon = faux_augure, faux_compagnon
    try:
        texte = get_contradictions([{"demande": "protéger ma mère"}], ["la fille s'y oppose"],
                                   "Le client demande un mandat de protection.",
                                   ["La mandante est apte", "La fille conteste"])
        donnees = json.loads(texte)
        verifier(isinstance(texte, str) and donnees["contradictions"], "get_contradictions rend du JSON")
        verifier("protéger ma mère" in demandes[-1] and "la fille s'y oppose" in demandes[-1],
                 "les demandes et commentaires sont envoyés")

        consequences = get_type_consequences([{"demande": "protéger ma mère"}], [],
                                             ["nullité", "inaptitude"], "Synthèse.",
                                             ["La mandante est apte"],
                                             droit_applicable=faux_compagnon("droit"))
        verifier(isinstance(consequences, list) and consequences[0]["id"] == "K1",
                 "get_type_consequences rend une liste")
        verifier(consequences[0]["source"] == {"titre": "", "url": "", "extrait": ""},
                 "une conséquence sans source est complétée par des champs vides")
        verifier("canlii.example" in demandes[-1], "les sources du Compagnon sont jointes")
    finally:
        analyse.ask_augure, analyse.ask_compagnon = analyse.ask_augure, analyse.ask_compagnon


def essai_acte_reel():
    """Produit un acte depuis un plan écrit à la main, sans appel à Augure."""
    print("\nacte complet (plan fourni, sans appel réseau)")
    for dossier in db.lister_dossiers():
        chemin = Path(create_acte_modifie(
            "Synthèse de démonstration.", [{"demande": "test"}], [], [{"fait": "élément"}], [],
            [], dossier=dossier,
            plan={"surlignages": [{"ancre": "mandataire", "note": "À vérifier."}],
                  "questions_au_client": ["Question de démonstration ?"]}))
        rapport = json.loads(chemin.with_suffix(".rapport.json").read_text(encoding="utf-8"))
        verifier(chemin.exists(), f"{dossier} : {chemin.name} écrit")
        verifier(rapport["fichier"] == str(chemin), f"{dossier} : le rapport pointe sur le fichier")


def main():
    sys.stdout.reconfigure(encoding="utf-8")
    # Une base jetable : les actes de démonstration n'encombrent pas data/documents.db.
    with tempfile.TemporaryDirectory() as temporaire:
        db.BASE, db.CACHE = Path(temporaire) / "essais.db", Path(temporaire) / "cache"
        for essai in (essai_json, essai_document, essai_word, essai_analyse, essai_acte_reel):
            essai()
    print(f"\n{_reussites} vérifications réussies, {_echecs} échecs")
    return 1 if _echecs else 0


if __name__ == "__main__":
    sys.exit(main())